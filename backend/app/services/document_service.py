"""Document ingestion and text extraction service supporting PDF, DOCX, TXT, and Markdown."""

import os
import re
import uuid
from pathlib import Path
from typing import List, Dict, Any, Tuple
import pymupdf
import docx
from app.core.config import settings
from app.core.exceptions import DocumentProcessingException
from app.core.logging import logger
from app.services.chunking_service import chunking_service

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
MAX_FILE_SIZE_BYTES = 25 * 1024 * 1024  # 25 MB limit


class DocumentService:
    """Manages document validation, extraction, and chunking pipeline."""

    def __init__(self):
        settings.ensure_directories()
        self.upload_dir = Path(settings.UPLOAD_DIR).resolve()

    def sanitize_filename(self, filename: str) -> str:
        """Strip directory paths and unsafe characters from filenames."""
        clean = os.path.basename(filename)
        clean = re.sub(r"[^\w\s.-]", "_", clean).strip()
        return clean or f"doc_{uuid.uuid4().hex[:8]}.txt"

    def validate_file(self, filename: str, file_size: int) -> str:
        """Validate file extension and file size."""
        ext = Path(filename).suffix.lower()
        if ext not in ALLOWED_EXTENSIONS:
            raise DocumentProcessingException(
                filename,
                f"Unsupported format '{ext}'. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}",
            )
        if file_size > MAX_FILE_SIZE_BYTES:
            raise DocumentProcessingException(
                filename,
                f"File size {file_size} bytes exceeds maximum allowed limit of {MAX_FILE_SIZE_BYTES} bytes (25MB)",
            )
        return ext

    def save_uploaded_file(self, filename: str, content: bytes) -> Tuple[str, str, int]:
        """Save file bytes to the local upload directory with a unique prefix."""
        clean_name = self.sanitize_filename(filename)
        unique_name = f"{uuid.uuid4().hex[:8]}_{clean_name}"
        destination = self.upload_dir / unique_name
        destination.parent.mkdir(parents=True, exist_ok=True)

        with open(destination, "wb") as f:
            f.write(content)

        return str(destination), clean_name, len(content)

    def extract_text_and_chunk(
        self,
        file_path: str,
        chunk_size: int = None,
        chunk_overlap: int = None,
    ) -> List[Dict[str, Any]]:
        """Extract text from file and chunk with page metadata."""
        p = Path(file_path)
        ext = p.suffix.lower()

        if ext == ".pdf":
            return self._extract_pdf(file_path, chunk_size, chunk_overlap)
        elif ext == ".docx":
            return self._extract_docx(file_path, chunk_size, chunk_overlap)
        elif ext in {".txt", ".md"}:
            return self._extract_text_file(file_path, chunk_size, chunk_overlap)
        else:
            raise DocumentProcessingException(p.name, f"Unsupported file extension {ext}")

    def _extract_pdf(
        self,
        file_path: str,
        chunk_size: int = None,
        chunk_overlap: int = None,
    ) -> List[Dict[str, Any]]:
        """Extract text from PDF pages using PyMuPDF."""
        chunks: List[Dict[str, Any]] = []
        try:
            doc = pymupdf.open(file_path)
            for page_idx in range(len(doc)):
                page = doc[page_idx]
                page_text = page.get_text("text").strip()
                if not page_text:
                    continue
                # Page number is 1-indexed for human readability
                page_chunks = chunking_service.split_text(
                    text=page_text,
                    chunk_size=chunk_size,
                    chunk_overlap=chunk_overlap,
                    page_number=page_idx + 1,
                )
                chunks.extend(page_chunks)
            doc.close()
            return chunks
        except Exception as e:
            logger.error(f"Failed to parse PDF {file_path}: {e}")
            raise DocumentProcessingException(Path(file_path).name, f"PDF extraction error: {str(e)}")

    def _extract_docx(
        self,
        file_path: str,
        chunk_size: int = None,
        chunk_overlap: int = None,
    ) -> List[Dict[str, Any]]:
        """Extract text from DOCX using python-docx."""
        try:
            doc = docx.Document(file_path)
            text_parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        text_parts.append(row_text)

            full_text = "\n\n".join(text_parts)
            return chunking_service.split_text(
                text=full_text,
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
                page_number=None,
            )
        except Exception as e:
            logger.error(f"Failed to parse DOCX {file_path}: {e}")
            raise DocumentProcessingException(Path(file_path).name, f"DOCX extraction error: {str(e)}")

    def _extract_text_file(
        self,
        file_path: str,
        chunk_size: int = None,
        chunk_overlap: int = None,
    ) -> List[Dict[str, Any]]:
        """Extract text from plain text or Markdown files."""
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
            return chunking_service.split_text(
                text=content,
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
                page_number=None,
            )
        except Exception as e:
            logger.error(f"Failed to parse text file {file_path}: {e}")
            raise DocumentProcessingException(Path(file_path).name, f"Text read error: {str(e)}")


document_service = DocumentService()
